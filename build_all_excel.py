#!/usr/bin/env python3
"""
build_all_excel.py — Buat Excel SEMUA lahan "Harga Sewa = 0" (seluruh Indonesia).

Kolom: No | No SPBU | Kode Lahan | Status Sewa
Membaca cache /tmp/price0_lahan.jsonl (hasil fetch), atau fetch ulang bila belum ada.

Contoh: python3 build_all_excel.py
"""
import os, re, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import scrapekit as sk
from sources import brightspace as bs

PAT = re.compile(r'^\d{7,9}-(IN|OU)-\d{1,4}$')


def main():
    items = bs.load_items()
    if not items:
        # fetch ulang bila cache kosong
        raw = sk.fetch_items_pages(bs.LAHAN_ITEMS_URL,
                                   [bs.price0_params(*j) for j in bs.all_pages()])
        for it in raw:
            items.setdefault(it['landCode'], it)

    rows = []
    for c in items:
        clean = re.sub(r'\s+', '', c)
        spbu = clean.split('-')[0]
        status = items[c].get('rentStatName') or ''
        rows.append((clean, spbu, status))
    rows.sort()

    bad = [c for c, _, _ in rows if not PAT.match(c)]
    print(f"Lahan unik: {len(rows)} | format aneh: {len(bad)}", bad[:8] if bad else "")

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Lahan Harga Sewa 0'
    cols = ['No', 'No SPBU', 'Kode Lahan', 'Status Sewa']
    sk.init_ws(ws, cols)
    for i, (code, spbu, status) in enumerate(rows, 1):
        r = i + 1
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=spbu)
        ws.cell(row=r, column=3, value=code)
        ws.cell(row=r, column=4, value=status)
    sk.style_ws(ws, [('A', 8), ('B', 12), ('C', 22), ('D', 13)],
                filter_ref=f"A1:D{len(rows)+1}")

    out = os.path.join(bs.DATA, 'lahan-hargasewa-0-all.xlsx')
    wb.save(out)
    print(f"OK: {out} | {len(rows)} baris | {os.path.getsize(out)/1024:.0f} KB")


if __name__ == '__main__':
    main()

#!/usr/bin/env python3
"""
embed_all_photos.py — Excel SEMUA lahan Harga Sewa=0 + thumbnail foto.

Kolom: No | No SPBU | Kode Lahan | Status Sewa | Foto
Membaca data/photo_map.json (hasil download_all_photos.py).

Contoh: python3 embed_all_photos.py
"""
import json, os, re, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import scrapekit as sk
from sources import brightspace as bs
from openpyxl import Workbook

THUMB_H = 84
COLS = ['No', 'No SPBU', 'Kode Lahan', 'Status Sewa', 'Foto']


def main():
    items = bs.load_items()
    url_map = json.load(open(bs.PHOTO_MAP)) if os.path.exists(bs.PHOTO_MAP) else {}

    rows = []
    for c, it in items.items():
        clean = re.sub(r'\s+', '', c)
        rows.append((clean, clean.split('-')[0], it.get('rentStatName') or '',
                     it.get('imageURL') or ''))
    rows.sort()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Lahan Harga Sewa 0'
    sk.init_ws(ws, COLS)

    no_photo = 0
    for i, (code, spbu, status, url) in enumerate(rows, 1):
        r = i + 1
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=spbu)
        ws.cell(row=r, column=3, value=code)
        ws.cell(row=r, column=4, value=status)
        rel = url_map.get(url)
        if rel and os.path.exists(os.path.join(bs.DATA, rel)):
            if not sk.embed_photo(ws, f'E{r}', os.path.join(bs.DATA, rel), THUMB_H):
                no_photo += 1
        else:
            no_photo += 1
        ws.row_dimensions[r].height = int(THUMB_H * 72 / 96) + 4

    sk.style_ws(ws, [('A', 8), ('B', 12), ('C', 22), ('D', 13), ('E', 13)],
                filter_ref=f"A1:E{len(rows)+1}")
    out = os.path.join(bs.DATA, 'lahan-hargasewa-0-all.xlsx')
    wb.save(out)
    print(f"OK: {out} | {len(rows)} baris, {len(rows)-no_photo} foto | "
          f"{os.path.getsize(out)/1024/1024:.1f} MB")


if __name__ == '__main__':
    main()

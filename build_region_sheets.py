#!/usr/bin/env python3
"""
build_region_sheets.py — Excel lahan Harga Sewa=0, dipisah per SHEET regional
(dalam satu workbook).

Sheet berdasarkan digit pertama No SPBU:
  1=Sumbagut 2=Sumbagsel 3=JBB 4=JBT 5=Jatimbalinus 6=Kalimantan 7=Sulawesi 8=Maluku Papua

Kolom tiap sheet: No | No SPBU | Kode Lahan | Wilayah | Status Sewa | Foto

Catatan: workbook dengan ~10.000 gambar bisa membuat Excel tidak merender
foto pada sheet akhir (limit ~7.500 gambar). Untuk semua foto tampil,
pakai build_region_files.py (8 file terpisah).

Contoh: python3 build_region_sheets.py
"""
import json, os, re, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import scrapekit as sk
from sources import brightspace as bs
from openpyxl import Workbook

THUMB_H = 84
COLS = ['No', 'No SPBU', 'Kode Lahan', 'Wilayah', 'Status Sewa', 'Foto']
OUT = os.path.join(bs.DATA, 'lahan-hargasewa-0-all-by-region.xlsx')


def main():
    items = bs.load_items()
    wilayah = bs.load_wilayah()
    url_map = json.load(open(bs.PHOTO_MAP)) if os.path.exists(bs.PHOTO_MAP) else {}

    wb = Workbook()
    wb.remove(wb.active)
    total_no_photo = 0

    for name in bs.REGIONS:
        rows = []
        for c, it in items.items():
            clean = re.sub(r'\s+', '', c)
            spbu = clean.split('-')[0]
            if bs.region_for(spbu) != name:
                continue
            rows.append((clean, spbu, wilayah.get(clean, ''),
                         it.get('rentStatName') or '', it.get('imageURL') or ''))
        rows.sort()

        ws = wb.create_sheet(title=name)
        sk.init_ws(ws, COLS)
        for i, (code, spbu, wil, status, url) in enumerate(rows, 1):
            r = i + 1
            ws.cell(row=r, column=1, value=i)
            ws.cell(row=r, column=2, value=spbu)
            ws.cell(row=r, column=3, value=code)
            ws.cell(row=r, column=4, value=wil)
            ws.cell(row=r, column=5, value=status)
            rel = url_map.get(url)
            if rel and os.path.exists(os.path.join(bs.DATA, rel)):
                if not sk.embed_photo(ws, f'F{r}', os.path.join(bs.DATA, rel), THUMB_H):
                    total_no_photo += 1
            else:
                total_no_photo += 1
            ws.row_dimensions[r].height = int(THUMB_H * 72 / 96) + 4
        sk.style_ws(ws, [('A', 7), ('B', 12), ('C', 22), ('D', 24), ('E', 13), ('F', 12)],
                    filter_ref=f"A1:F{len(rows)+1}")
        print(f"  {name}: {len(rows)} baris", flush=True)

    wb.save(OUT)
    print(f"OK: {OUT} | {os.path.getsize(OUT)/1024/1024:.1f} MB | tanpa foto: {total_no_photo}")


if __name__ == '__main__':
    main()

#!/usr/bin/env python3
"""
embed_photos.py — Sisipkan foto (thumbnail) ke Excel Jakarta Barat.

Menulis data/jakarta-barat-hargasewa-0-full.xlsx dengan kolom "Foto"
+ semua kolom data (sama seperti extract_lahan).

Contoh: python3 embed_photos.py
"""
import glob, json, os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import scrapekit as sk
from sources import brightspace as bs
from openpyxl import Workbook

THUMB_H = 100
COLS = ['Foto', 'No SPBU', 'Kode Lahan', 'Status Sewa', 'Harga Sewa', 'Periode',
        'Posisi', 'Luas (m2)', 'Panjang (m)', 'Lebar (m)', 'Daya Listrik (VA/W)',
        'Alamat', 'Lintang', 'Bujur']
POS = {'IN': 'Ruangan/Bangunan', 'OU': 'Tanah Kosong'}


def main():
    items = {c: it for c, it in bs.load_items().items() if bs.is_jakarta_barat(it)}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Lahan Harga 0 Jakarta Barat'
    sk.init_ws(ws, COLS)

    no_photo = 0
    for ri, (code, it) in enumerate(sorted(items.items()), start=2):
        ws.cell(row=ri, column=2, value=code.split('-')[0])
        ws.cell(row=ri, column=3, value=code)
        ws.cell(row=ri, column=4, value=it.get('rentStatName'))
        ws.cell(row=ri, column=5, value=it.get('rentalPrice'))
        ws.cell(row=ri, column=6, value=it.get('rentalPeriodName'))
        ws.cell(row=ri, column=7, value=f"{it.get('landPosition')} ({POS.get(it.get('landPosition'), '')})")
        ws.cell(row=ri, column=8, value=it.get('landArea'))
        ws.cell(row=ri, column=9, value=it.get('landLength'))
        ws.cell(row=ri, column=10, value=it.get('landWidth'))
        ws.cell(row=ri, column=11, value=it.get('powerCapacity'))
        ws.cell(row=ri, column=12, value=it.get('address'))
        ws.cell(row=ri, column=13, value=it.get('latitude'))
        ws.cell(row=ri, column=14, value=it.get('longitude'))

        found = glob.glob(os.path.join(bs.DATA, 'photos', code, 'img-1.*'))
        if found:
            if not sk.embed_photo(ws, f'A{ri}', found[0], THUMB_H):
                no_photo += 1
        else:
            no_photo += 1
        ws.row_dimensions[ri].height = int(THUMB_H * 72 / 96) + 6

    sk.style_ws(ws, [('A', 16), ('B', 10), ('C', 20), ('D', 11), ('E', 11), ('F', 9),
                     ('G', 22), ('H', 10), ('I', 11), ('J', 10), ('K', 16), ('L', 55),
                     ('M', 10), ('N', 10)])
    out = os.path.join(bs.DATA, 'jakarta-barat-hargasewa-0-full.xlsx')
    wb.save(out)
    print(f"OK: {out} | {len(items)} baris, {len(items)-no_photo} foto, tanpa foto {no_photo} "
          f"| {os.path.getsize(out)/1024/1024:.1f} MB")


if __name__ == '__main__':
    main()

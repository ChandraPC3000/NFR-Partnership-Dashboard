#!/usr/bin/env python3
"""
scrapekit.py — Helper umum untuk scraping + output Excel/CSV.

Bagian yang bisa dipakai ulang di semua sumber data (tidak spesifik
Brightspace): HTTP dengan retry, pagination, cache JSONL, thumbnail foto,
dan penulisan Excel (header, gambar, filter).

Pola pemakaian:
    import scrapekit as sk
    items = sk.load_items_from_jsonl(path)
    sk.init_ws(ws, cols); sk.style_ws(ws, widths, filter_ref=...)
    ok = sk.embed_photo(ws, 'F2', path)
"""
import json, os, sys, time, io
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed
from PIL import Image as PILImage
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)',
    'X-Requested-With': 'XMLHttpRequest',
    'Accept': 'application/json',
}
NAVY = '0D2F4F'


# ---------------------------------------------------------------- HTTP
def session():
    s = requests.Session()
    s.headers.update(HEADERS)
    return s


def fetch_json(url, params=None, retries=5, timeout=60):
    """GET JSON dengan retry. Kembalikan dict, atau None bila gagal."""
    for a in range(retries):
        try:
            r = requests.get(url, params=params, timeout=timeout, headers=HEADERS)
            if r.status_code == 200 and r.content:
                return r.json()
        except Exception as e:
            sys.stderr.write(f"  retry {a+1} {url.split('/')[-1]}: {e}\n")
        time.sleep(2 + 1.5 * a)
    return None


def fetch_items_pages(url, jobs, max_workers=4, delay=0.15, retries=5, timeout=60):
    """Ambil beberapa halaman (jobs = list param dict) -> daftar item gabungan."""
    def one(job):
        d = fetch_json(url, params=job, retries=retries, timeout=timeout)
        return (d or {}).get('items') or []
    raw = []
    with ThreadPoolExecutor(max_workers=max_workers) as ex:
        for fut in as_completed([ex.submit(one, j) for j in jobs]):
            raw.extend(fut.result())
            time.sleep(delay)
    return raw


def download_image(url, path, retries=4, timeout=60):
    """Unduh gambar ke path (tanpa ekstensi). Kembalikan (True, path.ext, size)
    atau (False, path, 0)."""
    def ext_for(ct):
        ct = (ct or '').lower()
        if 'png' in ct: return 'png'
        if 'webp' in ct: return 'webp'
        if 'gif' in ct: return 'gif'
        return 'jpg'
    for a in range(retries):
        try:
            r = requests.get(url, timeout=timeout,
                             headers={'User-Agent': 'Mozilla/5.0', 'Accept': 'image/*,*/*;q=0.8'})
            if r.status_code == 200 and r.content and not r.content[:1] == b'{':
                final = f"{path}.{ext_for(r.headers.get('Content-Type'))}"
                with open(final, 'wb') as f:
                    f.write(r.content)
                return True, final, len(r.content)
            elif r.status_code == 200:
                return False, path, 0
        except Exception as e:
            sys.stderr.write(f"  retry {os.path.basename(path)}: {e}\n")
        time.sleep(1.5)
    return False, path, 0


def image_urls(it):
    """URL foto dari satu item (imageURL + imageURLs), dedup, urut terjaga."""
    import re
    urls = []
    u = it.get('imageURL') or ''
    if u:
        urls.append(u)
    ul = it.get('imageURLs') or []
    if isinstance(ul, list):
        urls.extend(ul)
    elif isinstance(ul, str):
        urls.extend(re.findall(r'https?://[^\s"\\]+', ul))
    seen, out = set(), []
    for x in urls:
        if x and x not in seen:
            seen.add(x); out.append(x)
    return out


# ---------------------------------------------------------------- cache
def load_items_from_jsonl(path, key='landCode'):
    """Baca cache JSONL (tiap baris {items:[...]}) -> dict keyed by `key`."""
    items = {}
    if os.path.exists(path):
        for line in open(path):
            try:
                r = json.loads(line)
                for it in r.get('items', []):
                    items.setdefault(it[key], it)
            except Exception:
                pass
    return items


# ---------------------------------------------------------------- Excel
def init_ws(ws, cols):
    """Tulis header navy di baris 1."""
    hf = PatternFill('solid', fgColor=NAVY)
    hfont = Font(color='FFFFFF', bold=True)
    for ci, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=ci, value=c)
        cell.fill = hf
        cell.font = hfont


def style_ws(ws, widths, freeze='A2', filter_ref=None):
    """Lebar kolom (list (huruf, lebar)), freeze, dan autofilter."""
    for col, w in widths:
        ws.column_dimensions[col].width = w
    ws.freeze_panes = freeze
    if filter_ref:
        ws.auto_filter.ref = filter_ref


def make_thumb(path, max_h=84, max_w=360, quality=72):
    """Buat thumbnail JPEG di memori -> (BytesIO, width, height)."""
    im = PILImage.open(path).convert('RGB')
    im.thumbnail((max_w, max_h), PILImage.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, 'JPEG', quality=quality)
    buf.seek(0)
    return buf, im.width, im.height


def embed_photo(ws, cell, filepath, thumb_h=84):
    """Sisipkan thumbnail ke sel. Kembalikan True bila berhasil."""
    try:
        buf, w, h = make_thumb(filepath, max_h=thumb_h)
        xl = XLImage(buf)
        xl.width, xl.height = w, h
        ws.add_image(xl, cell)
        return True
    except Exception:
        return False
